#!/usr/bin/env python3
"""Build AHRI go-live keep-lists from reference zip or extracted reference folder."""

from __future__ import annotations

import argparse
import csv
import re
import sys
import zipfile
from collections import Counter
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    print("openpyxl is required: pip install openpyxl", file=sys.stderr)
    sys.exit(1)

REFERENCE_FILES = {
    "sampledata": "AHRI/Biorepository/sampledata.xlsx",
    "rooms": "AHRI/Biorepository/Room and freezer detail.xlsx",
    "equipment": "AHRI/All-EquipmentList.xlsx",
}


def normalize_name(value: object) -> str:
    if value is None:
        return ""
    text = str(value).strip().lower()
    text = re.sub(r"\s+", " ", text)
    return text


def resolve_input_path(source: Path) -> dict[str, Path]:
    if source.is_file() and source.suffix.lower() == ".zip":
        extracted: dict[str, Path] = {}
        with zipfile.ZipFile(source) as archive:
            for key, member in REFERENCE_FILES.items():
                if member not in archive.namelist():
                    print(f"Warning: missing {member} in zip", file=sys.stderr)
                    continue
                target = Path("/tmp/ahri-keep-lists") / Path(member).name
                target.parent.mkdir(parents=True, exist_ok=True)
                target.write_bytes(archive.read(member))
                extracted[key] = target
        return extracted

    base = source
    return {
        "sampledata": base / "Biorepository" / "sampledata.xlsx",
        "rooms": base / "Biorepository" / "Room and freezer detail.xlsx",
        "equipment": base / "All-EquipmentList.xlsx",
    }


def write_barcodes(sampledata_path: Path, output_dir: Path) -> int:
    wb = load_workbook(sampledata_path, read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    headers = [normalize_name(c).replace(" ", "_") for c in next(rows)]
    try:
        sno_idx = headers.index("s_no")
        barcode_idx = headers.index("sample_id")
    except ValueError as exc:
        raise SystemExit(f"Unexpected sampledata headers: {headers}") from exc

    out = output_dir / "biorepository-barcodes.csv"
    count = 0
    with out.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["barcode", "manifest_sno", "freezer_no", "normalized_barcode"])
        for row in rows:
            barcode = row[barcode_idx] if barcode_idx < len(row) else None
            if not barcode:
                continue
            sno = row[sno_idx] if sno_idx < len(row) else None
            freezer = row[headers.index("freezer_no")] if "freezer_no" in headers else ""
            writer.writerow([str(barcode).strip(), sno, freezer, normalize_name(barcode)])
            count += 1
    wb.close()
    return count


def write_storage_hints(sampledata_path: Path, output_dir: Path) -> int:
    wb = load_workbook(sampledata_path, read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    headers = [normalize_name(c).replace(" ", "_") for c in next(rows)]

    def col(name: str) -> int | None:
        return headers.index(name) if name in headers else None

    freezer_idx = col("freezer_no")
    shelf_idx = col("shelf_no.")
    rack_idx = col("rack_no.")
    box_idx = col("box_no.")
    zone_idx = col("zone")

    combos: Counter[tuple[str, str, str, str, str]] = Counter()
    for row in rows:
        values = []
        for idx in (zone_idx, freezer_idx, shelf_idx, rack_idx, box_idx):
            if idx is None or idx >= len(row) or row[idx] in (None, ""):
                values.append("")
            else:
                values.append(str(row[idx]).strip())
        key = tuple(values)  # type: ignore[assignment]
        if any(key):
            combos[key] += 1

    out = output_dir / "storage-hierarchy-hints.csv"
    with out.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["zone", "freezer", "shelf", "rack", "box", "sample_count"])
        for key, sample_count in sorted(combos.items(), key=lambda item: item[0]):
            writer.writerow([*key, sample_count])
    wb.close()
    return len(combos)


def write_rooms(rooms_path: Path, output_dir: Path) -> int:
    wb = load_workbook(rooms_path, read_only=True, data_only=True)
    ws = wb.active
    out = output_dir / "storage-rooms-freezers.csv"
    rows_written = 0
    with out.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["item", "quantity", "normalized_item"])
        for row in ws.iter_rows(min_row=2, values_only=True):
            item = row[1] if len(row) > 1 else None
            qty = row[2] if len(row) > 2 else None
            if not item or not str(item).strip():
                continue
            item_text = str(item).strip()
            if item_text.lower() in {"item", "storage infrastructure detail:"}:
                continue
            writer.writerow([item_text, qty, normalize_name(item_text)])
            rows_written += 1
    wb.close()
    return rows_written


def write_equipment(equipment_path: Path, output_dir: Path) -> int:
    if not equipment_path.exists():
        return 0
    wb = load_workbook(equipment_path, read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    headers = [normalize_name(c) for c in next(rows)]
    name_idx = 0
    for idx, header in enumerate(headers):
        if header in {"equipment", "equipment_name", "name", "item"}:
            name_idx = idx
            break

    out = output_dir / "equipment-keep.csv"
    count = 0
    with out.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["equipment_name", "normalized_name"])
        for row in rows:
            if not row or name_idx >= len(row):
                continue
            name = row[name_idx]
            if not name or not str(name).strip():
                continue
            text = str(name).strip()
            writer.writerow([text, normalize_name(text)])
            count += 1
    wb.close()
    return count


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "source",
        nargs="?",
        default="/home/dawit/Downloads/AHRI-20260525T055818Z-3-001.zip",
        help="Reference zip or extracted AHRI/ folder",
    )
    parser.add_argument(
        "-o",
        "--output-dir",
        default="scripts/ahri/reference",
        help="Directory for generated CSV keep-lists",
    )
    args = parser.parse_args()

    source = Path(args.source)
    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    paths = resolve_input_path(source)
    barcode_count = write_barcodes(paths["sampledata"], output_dir)
    storage_count = write_storage_hints(paths["sampledata"], output_dir)
    room_count = write_rooms(paths["rooms"], output_dir)
    equipment_count = write_equipment(paths.get("equipment", Path()), output_dir)

    print(f"Wrote keep-lists to {output_dir}")
    print(f"  biorepository barcodes: {barcode_count}")
    print(f"  storage hierarchy hints: {storage_count}")
    print(f"  room/freezer rows: {room_count}")
    print(f"  equipment rows: {equipment_count}")


if __name__ == "__main__":
    main()
