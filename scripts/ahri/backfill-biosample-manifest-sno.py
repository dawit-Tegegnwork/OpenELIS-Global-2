#!/usr/bin/env python3
"""Extract barcode -> manifest_sno mapping from AHRI sampledata.xlsx for DB backfill."""

from __future__ import annotations

import argparse
import csv
import re
import sys
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    print("openpyxl is required: pip install openpyxl", file=sys.stderr)
    sys.exit(1)

SNO_HEADERS = {"s_no", "sno", "s.no", "serial", "serial_no", "manifest_sno"}
BARCODE_HEADERS = {"sample_id", "sampleid", "barcode", "sample id"}


def normalize_header(value: object) -> str:
    if value is None:
        return ""
    text = str(value).strip().lower()
    text = re.sub(r"[\s._-]+", "_", text)
    return text.strip("_")


def pick_column(headers: list[str], candidates: set[str]) -> int | None:
    for idx, header in enumerate(headers):
        if header in candidates:
            return idx
    return None


def coerce_sno(value: object) -> int | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return int(value)
    text = str(value).strip()
    if not text:
        return None
    try:
        return int(float(text))
    except ValueError:
        return None


def coerce_barcode(value: object) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def extract_mapping(xlsx_path: Path) -> list[tuple[str, int]]:
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows)
    except StopIteration:
        raise SystemExit(f"No rows found in {xlsx_path}")

    headers = [normalize_header(cell) for cell in header_row]
    sno_idx = pick_column(headers, SNO_HEADERS)
    barcode_idx = pick_column(headers, BARCODE_HEADERS)
    if sno_idx is None or barcode_idx is None:
        raise SystemExit(
            f"Could not find S.No and Barcode columns in {xlsx_path}. Headers: {headers}"
        )

    by_key: dict[str, tuple[str, int]] = {}
    for row in rows:
        barcode = coerce_barcode(row[barcode_idx] if barcode_idx < len(row) else None)
        sno = coerce_sno(row[sno_idx] if sno_idx < len(row) else None)
        if not barcode or sno is None:
            continue
        key = barcode.lower()
        existing = by_key.get(key)
        if existing is None or sno > existing[1]:
            by_key[key] = (barcode, sno)

    wb.close()
    return sorted(by_key.values(), key=lambda item: item[1])


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "xlsx",
        nargs="?",
        default="scripts/ahri/reference/sampledata.xlsx",
        help="Path to sampledata.xlsx",
    )
    parser.add_argument(
        "-o",
        "--output",
        default="scripts/ahri/data/manifest-sno-backfill.csv",
        help="Output CSV path (barcode,manifest_sno)",
    )
    args = parser.parse_args()

    xlsx_path = Path(args.xlsx)
    output_path = Path(args.output)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    mapping = extract_mapping(xlsx_path)
    with output_path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["barcode", "manifest_sno"])
        writer.writerows(mapping)

    print(f"Wrote {len(mapping)} rows to {output_path}")


if __name__ == "__main__":
    main()
